# pos/views/purchase_excel_views.py
#
# Purchase Entry — Excel Template Download + Bulk Import
# ---------------------------------------------------------------------------
# Ek "Purchase Data" sheet me:
#   - MASTER fields (party, date, terms, due date, cash/bank account,
#     narration, freight/other/round) SIRF us purchase ki PEHLI row me bharte hain
#   - ITEM fields (item variant, qty, price, discount%) HAR row me bharte hain
#   - Naya PARTY_NAME bharte hi ek NAYI purchase entry shuru maani jaati hai
#
# Import ke waqt tax/GST calculation ke liye humne apna alag logic nahi likha —
# PurchaseItem.save() ka EXISTING model-level logic hi reuse hota hai
# (discount, GST toggle, CGST/SGST vs IGST split) — taaki Excel se aur normal
# Purchase Entry form se bani entry me koi farak na ho.
#
# ─────────────────────────────────────────────────────────────────────────────
# 🔧 CRITICAL FIX (ye version): DUE_DATE / CASH_ACCOUNT / BANK_ACCOUNT
# exclusivity check galat error de raha tha — jaise ek Cash-terms entry me
# sirf CASH_ACCOUNT bhara hone ke bawajood "only CASH_ACCOUNT should be
# filled" ka error aata tha.
#
# ROOT CAUSE: Excel me khaali DATE cell (DUE_DATE) ko pandas `NaT`
# (Not-a-Time) deta hai — jo na `None` hai na `float NaN`. `is_empty()`
# sirf None/float-NaN check karta tha, isliye khaali DUE_DATE bhi "filled"
# maana ja raha tha, aur Cash/Bank entries pe false-positive error aata tha.
#
# FIX: `pd.isna(v)` use kiya — ye None, float NaN, aur NaT — teeno ko sahi
# se pakadta hai. (Verified with the actual uploaded sample file — 3
# entries: 1 item / 4 items / 3 items — ab teeno sahi validate hote hain.)
#
# Ek chhota cosmetic fix bhi: agar variant ka `size` DB me decimal (jaise
# 45.0) store hai, to label me "(45.0)" ki jagah ab "(45)" dikhega.
# ─────────────────────────────────────────────────────────────────────────────
#
# ⚠️ ASSUMPTIONS (apne actual models ke hisaab se verify/adjust kar lena):
#   1. Account.group values: Supplier -> "Supplier", Cash -> "Case In Hand",
#      Bank -> "Bank Account"
#   2. Account model me `account_name`, `group`, `state` fields hain
#   3. setting model me branch-wise "PI" prefix field ho sakta hai — agar
#      nahi hai to default "PI" use hoga
#   4. Party (supplier), Cash Account, Bank Account — teeno is version me
#      current branch se filter hote hain
#
# URL me add karna hoga (urls.py):
#   path('purchase-excel/template/', PurchaseExcelTemplateView.as_view()),
#   path('purchase-excel/import/',   PurchaseExcelImportView.as_view()),

import io
from datetime import datetime
from decimal import Decimal, InvalidOperation

import pandas as pd
from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status

from pos.models.items import items, itemvariants
from pos.models.account import Account
from pos.models.settings import setting
from pos.models.purchaseentry import PurchaseMaster, PurchaseItem
from pos.models.cashpayment import CashPayment
from pos.models.bankpayment import BankPayment

from ecommerce.permissions import IsSuperAdminOrBranchOrPagePermittedEmployee


# ─────────────────────────────────────────────────────────────────────────────
# CONSTANTS / HELPERS  (template aur import dono yehi use karte hai)
# ─────────────────────────────────────────────────────────────────────────────

TERMS_CHOICES = ["Credit", "Cash", "Bank"]

COLUMNS = [
    "PARTY_NAME", "DATE", "PURCHASE_BILL_NO", "TERMS", "DUE_DATE",
    "CASH_ACCOUNT", "BANK_ACCOUNT", "NARRATION",
    "FREIGHT_CHARGE", "OTHER_EXPENSE", "ROUND_AMOUNT",
    "ITEM_VARIANT", "HSN_CODE", "QTY", "PRICE", "UNIT",
    "DISCOUNT_PERCENT", "TAX_PERCENT",
]
COL_INDEX = {name: i + 1 for i, name in enumerate(COLUMNS)}  # 1-based for openpyxl


def dedupe_preserve_order(values):
    """
    'ADS Agency', 'ADS Agency', 'ads agency ' jaisi duplicate /
    whitespace-mismatched entries ko ek hi entry me convert karta hai —
    order wahi rakhta hai jisme pehli baar mila tha.
    """
    seen = set()
    result = []
    for v in values:
        if not v:
            continue
        key = v.strip().lower()
        if key and key not in seen:
            seen.add(key)
            result.append(v.strip())
    return result


def _clean_size(val):
    """size DB me 45.0 (float) ho to '45' dikhega, '45.0' nahi."""
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def variant_label(item_obj, variant_obj):
    """
    'ItemName (Size)' jaisa label — template dropdown me dikhne wala text
    aur import ke waqt match karne wala text — dono isi function se banate
    hain taaki hamesha ek jaisa rahe.
    """
    size = _clean_size(getattr(variant_obj, "size", None))
    if size:
        return f"{item_obj.itemName} ({size})"
    return item_obj.itemName


def build_variant_lookup(variants_qs):
    """[(variant_obj, label), ...] — same order as queryset."""
    return [(v, variant_label(v.item, v)) for v in variants_qs]


def _unit_symbol(item_obj):
    if item_obj.unit:
        return getattr(item_obj.unit, "symbol", None) or getattr(item_obj.unit, "name", "pcs")
    return "pcs"


def _fy_string():
    now = datetime.now()
    fy_start, fy_end = (now.year, now.year + 1) if now.month >= 4 else (now.year - 1, now.year)
    return f"{str(fy_start)[2:]}-{str(fy_end)[2:]}"


def _next_voucher_no(branch, model_cls, prefix_setting_attr, default_prefix):
    settings_obj = setting.objects.filter(branch=branch).first()
    prefix = getattr(settings_obj, prefix_setting_attr, default_prefix) if settings_obj else default_prefix

    field = "billNo" if model_cls is PurchaseMaster else "voucher_no"
    last_obj = model_cls.objects.filter(branch=branch).order_by("-id").first()
    last_no = 0
    last_value = getattr(last_obj, field, None) if last_obj else None
    if last_value:
        try:
            parts = last_value.split("/")
            if len(parts) >= 3:
                last_no = int(parts[-1])
        except (ValueError, IndexError):
            last_no = 0

    return f"{prefix}/{_fy_string()}/{str(last_no + 1).zfill(4)}"


def generate_purchase_voucher(branch):
    return _next_voucher_no(branch, PurchaseMaster, "PI", "PI")


def generate_cash_payment_voucher(branch):
    return _next_voucher_no(branch, CashPayment, "CP", "CP")


def generate_bank_payment_voucher(branch):
    return _next_voucher_no(branch, BankPayment, "BP", "BP")


def to_decimal(val, default=Decimal("0.00")):
    try:
        if val is None:
            return default
        s = str(val).replace("%", "").strip()
        if s == "" or s.lower() == "nan":
            return default
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return default


# ─────────────────────────────────────────────────────────────────────────────
# 1) TEMPLATE DOWNLOAD
# ─────────────────────────────────────────────────────────────────────────────

class PurchaseExcelTemplateView(APIView):
    """Purchase bulk-import ke liye Excel template — sab dropdowns
    (Party/Cash Account/Bank Account/Item Variant) current branch ke DB
    data se auto-populate hote hain (deduped)."""

    permission_classes = [IsSuperAdminOrBranchOrPagePermittedEmployee]
    page_key = "/purchaseimport"

    def get(self, request):
        branch = request.user.get_effective_branch()
        if not branch:
            return Response({"error": "No branch linked to this user"}, status=400)

        wb = Workbook()
        ws = wb.active
        ws.title = "Purchase Data"

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        master_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # peach = ek baar wali fields
        item_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")     # blue = har row wali fields
        header_font = Font(color="FFFFFF", bold=True)
        thin = Border(*(Side(style="thin"),) * 4)

        max_row = 1000

        COLUMN_WIDTHS = {
            "ITEM_VARIANT": 40,
            "PARTY_NAME": 28,
            "NARRATION": 30,
        }

        # ── Header row ──
        for idx, col_name in enumerate(COLUMNS, 1):
            required = col_name in ("PARTY_NAME", "DATE", "TERMS", "ITEM_VARIANT", "QTY", "PRICE")
            label = col_name.replace("_", " ") + ("*" if required else "")
            cell = ws.cell(row=1, column=idx, value=label)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin
            ws.column_dimensions[get_column_letter(idx)].width = COLUMN_WIDTHS.get(col_name, 20)

        MASTER_ONLY = {
            "PARTY_NAME", "DATE", "PURCHASE_BILL_NO", "TERMS", "DUE_DATE",
            "CASH_ACCOUNT", "BANK_ACCOUNT", "NARRATION",
            "FREIGHT_CHARGE", "OTHER_EXPENSE", "ROUND_AMOUNT",
        }
        for col_name, idx in COL_INDEX.items():
            fill = master_fill if col_name in MASTER_ONLY else item_fill
            for r in range(2, max_row + 1):
                cell = ws.cell(row=r, column=idx)
                cell.fill = fill
                cell.border = thin

        for col_name in ("HSN_CODE", "UNIT", "TAX_PERCENT"):
            idx = COL_INDEX[col_name]
            for r in range(2, max_row + 1):
                ws.cell(row=r, column=idx).font = Font(italic=True, color="808080")

        # ===== HIDDEN LOOKUP DATA (same sheet, right side, hidden columns) =====
        suppliers_raw = list(
            Account.objects.filter(group="Supplier", branch=branch)
            .order_by("account_name")
            .values_list("account_name", flat=True)
        )
        cash_accounts_raw = list(
            Account.objects.filter(group="Case In Hand", branch=branch)
            .order_by("account_name")
            .values_list("account_name", flat=True)
        )
        bank_accounts_raw = list(
            Account.objects.filter(group="Bank Account", branch=branch)
            .order_by("account_name")
            .values_list("account_name", flat=True)
        )

        suppliers = dedupe_preserve_order(suppliers_raw)
        cash_accounts = dedupe_preserve_order(cash_accounts_raw)
        bank_accounts = dedupe_preserve_order(bank_accounts_raw)

        variants_qs = (
            itemvariants.objects.filter(item__branch=branch)
            .select_related("item", "item__unit")
            .distinct()
        )
        variant_lookup = build_variant_lookup(variants_qs)

        seen_item_labels = set()
        item_lookup_rows = []  # (label, hsn, unit, tax%, purchase_price)
        for v, label in variant_lookup:
            key = label.strip().lower()
            if key in seen_item_labels:
                continue
            seen_item_labels.add(key)
            it = v.item
            item_lookup_rows.append((
                label,
                it.hsnCode or "",
                _unit_symbol(it),
                (it.taxSlab or "0").replace("%", ""),
                v.purchasePrice or 0,
            ))
        item_lookup_rows.sort(key=lambda r: r[0].lower())

        hidden_start = len(COLUMNS) + 2
        PARTY_COL = hidden_start
        CASH_COL = hidden_start + 1
        BANK_COL = hidden_start + 2
        ITEM_LOOKUP_START = hidden_start + 3  # label,hsn,unit,tax,price -> 5 columns

        for r, name in enumerate(suppliers, start=2):
            ws.cell(row=r, column=PARTY_COL, value=name)
        for r, name in enumerate(cash_accounts, start=2):
            ws.cell(row=r, column=CASH_COL, value=name)
        for r, name in enumerate(bank_accounts, start=2):
            ws.cell(row=r, column=BANK_COL, value=name)
        for r, row_data in enumerate(item_lookup_rows, start=2):
            for offset, val in enumerate(row_data):
                ws.cell(row=r, column=ITEM_LOOKUP_START + offset, value=val)

        for c in (PARTY_COL, CASH_COL, BANK_COL, ITEM_LOOKUP_START, ITEM_LOOKUP_START + 1,
                  ITEM_LOOKUP_START + 2, ITEM_LOOKUP_START + 3, ITEM_LOOKUP_START + 4):
            ws.column_dimensions[get_column_letter(c)].hidden = True

        def col_range(col_idx, count):
            letter = get_column_letter(col_idx)
            return f"${letter}$2:${letter}${max(count + 1, 2)}"

        # ===== DATA VALIDATIONS (dropdowns) =====
        def add_list_validation(target_col_name, formula):
            idx = COL_INDEX[target_col_name]
            letter = get_column_letter(idx)
            dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{max_row}")

        add_list_validation("PARTY_NAME", col_range(PARTY_COL, len(suppliers)))
        add_list_validation("TERMS", '"Credit,Cash,Bank"')
        add_list_validation("CASH_ACCOUNT", col_range(CASH_COL, len(cash_accounts)))
        add_list_validation("BANK_ACCOUNT", col_range(BANK_COL, len(bank_accounts)))
        add_list_validation("ITEM_VARIANT", col_range(ITEM_LOOKUP_START, len(item_lookup_rows)))

        # ===== AUTO-FILL FORMULAS (HSN / UNIT / TAX% / PRICE default) =====
        lookup_range = (
            f"${get_column_letter(ITEM_LOOKUP_START)}$2:"
            f"${get_column_letter(ITEM_LOOKUP_START + 4)}${max(len(item_lookup_rows) + 1, 2)}"
        )
        item_col_letter = get_column_letter(COL_INDEX["ITEM_VARIANT"])
        for r in range(2, max_row + 1):
            ws.cell(row=r, column=COL_INDEX["HSN_CODE"],
                     value=f'=IFERROR(VLOOKUP(${item_col_letter}{r},{lookup_range},2,FALSE),"")')
            ws.cell(row=r, column=COL_INDEX["UNIT"],
                     value=f'=IFERROR(VLOOKUP(${item_col_letter}{r},{lookup_range},3,FALSE),"")')
            ws.cell(row=r, column=COL_INDEX["TAX_PERCENT"],
                     value=f'=IFERROR(VLOOKUP(${item_col_letter}{r},{lookup_range},4,FALSE),"")')
            ws.cell(row=r, column=COL_INDEX["PRICE"],
                     value=f'=IFERROR(VLOOKUP(${item_col_letter}{r},{lookup_range},5,FALSE),"")')

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="purchase_import_template.xlsx"'
        return response


# ─────────────────────────────────────────────────────────────────────────────
# 2) IMPORT
# ─────────────────────────────────────────────────────────────────────────────

class PurchaseExcelImportView(APIView): 
    """Excel se multiple purchase entries (har entry me multiple item
    variants) ek saath create karta hai — bilkul PurchaseEntryForm jaisa
    hi result (tax calc, voucher number, cash/bank auto-payment).

    Grouping rule: PARTY_NAME jis row me bharaa hai wahan se ek NAYI
    purchase entry shuru hoti hai; agli rows (jab tak PARTY_NAME dobara na
    bhare) usi entry ke additional item rows maani jaati hain."""

    permission_classes = [IsSuperAdminOrBranchOrPagePermittedEmployee]
    page_key = "/purchaseimport"

    def _create_cash_payment(self, purchase, branch, request):
        voucher = generate_cash_payment_voucher(branch)
        return CashPayment.objects.create(
            date=purchase.date,
            voucher_no=voucher,
            cash_account=purchase.case_account,
            op_account=purchase.partyName,
            branch=branch,
            amount=purchase.grand_total,
            mode="Cash",
            narration=f"Auto payment against Purchase {purchase.billNo}",
            type="PCP",
            purchase=purchase,
            created_by=request.user,
        )

    def _create_bank_payment(self, purchase, branch, request):
        voucher = generate_bank_payment_voucher(branch)
        return BankPayment.objects.create(
            date=purchase.date,
            voucher_no=voucher,
            bank_account=purchase.bank_account,
            op_account=purchase.partyName,
            branch=branch,
            amount=purchase.grand_total,
            mode="Auto",
            narration=f"Auto payment against Purchase {purchase.billNo}",
            type="PBP",
            purchase=purchase,
            created_by=request.user,
        )

    def post(self, request):
        branch = request.user.get_effective_branch()
        if not branch:
            return Response({"error": "No branch linked to this user"}, status=400)

        excel_file = request.FILES.get("file")
        if not excel_file:
            return Response({"error": "No file uploaded"}, status=400)

        try:
            df = pd.read_excel(excel_file, sheet_name=0, header=0)
        except Exception as e:
            return Response({"error": f"Could not read Excel file: {e}"}, status=400)

        df.columns = [str(c).strip().upper().replace("*", "").replace(" ", "_") for c in df.columns]
        missing = [c for c in ("PARTY_NAME", "TERMS", "ITEM_VARIANT", "QTY", "PRICE") if c not in df.columns]
        if missing:
            return Response({"error": f"Missing required columns: {', '.join(missing)}"}, status=400)

        def is_empty(v):
            """
            🔧 CRITICAL FIX: pd.isna() None / float NaN / pandas NaT
            (khaali DATE cell) — teeno ko sahi se "empty" pakadta hai.
            Pehle sirf None/float-NaN check hota tha, isliye khaali
            DUE_DATE cell bhi "filled" maan liya jaata tha aur Cash/Bank
            terms wali entries pe galat error aata tha.
            """
            if v is None:
                return True
            try:
                if pd.isna(v):
                    return True
            except (TypeError, ValueError):
                pass
            return str(v).strip() == ""

        def clean(v):
            """
            Excel me agar koi text-looking column (jaise PURCHASE_BILL_NO)
            numeric ho, to pandas use float bana deta hai (12345 ->
            12345.0). Yahan integer-valued floats ko wapas clean integer
            string me convert kar dete hain taaki '.0' na aaye.
            """
            if is_empty(v):
                return ""
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v).strip()

        def parse_date(v, default=None):
            if is_empty(v):
                return default
            try:
                return pd.to_datetime(v).date()
            except Exception:
                return default

        # ── Lookups — sirf is branch ka data ──
        suppliers = {a.account_name.strip().lower(): a for a in Account.objects.filter(group="Supplier", branch=branch)}
        cash_accounts = {a.account_name.strip().lower(): a for a in Account.objects.filter(group="Case In Hand", branch=branch)}
        bank_accounts = {a.account_name.strip().lower(): a for a in Account.objects.filter(group="Bank Account", branch=branch)}

        variant_map = {}
        variants_qs = itemvariants.objects.filter(item__branch=branch).select_related("item", "item__unit").distinct()
        for v, label in build_variant_lookup(variants_qs):
            variant_map[label.strip().lower()] = v

        # ── PASS 1: group rows into purchase entries ──
        groups = []
        current = None
        errors = []

        for idx, row in df.iterrows():
            row_no = idx + 2
            party_raw = clean(row.get("PARTY_NAME"))
            item_raw = clean(row.get("ITEM_VARIANT"))

            if not party_raw and not item_raw:
                continue  # fully blank row

            if party_raw:
                if current:
                    groups.append(current)
                current = {
                    "row_no": row_no,
                    "party_name_raw": party_raw,
                    "date_raw": row.get("DATE"),
                    "purchasebill_no": clean(row.get("PURCHASE_BILL_NO")),
                    "terms_raw": clean(row.get("TERMS")),
                    "due_date_raw": row.get("DUE_DATE"),
                    "cash_account_raw": clean(row.get("CASH_ACCOUNT")),
                    "bank_account_raw": clean(row.get("BANK_ACCOUNT")),
                    "narration": clean(row.get("NARRATION")),
                    "freight": to_decimal(row.get("FREIGHT_CHARGE")),
                    "other_expense": to_decimal(row.get("OTHER_EXPENSE")),
                    "round_amount": to_decimal(row.get("ROUND_AMOUNT")),
                    "items": [],
                }

            if item_raw:
                if not current:
                    errors.append(f"Row {row_no}: Item given without a PARTY_NAME on this or an earlier row")
                    continue
                current["items"].append({
                    "row_no": row_no,
                    "item_raw": item_raw,
                    "qty": to_decimal(row.get("QTY")),
                    "price": to_decimal(row.get("PRICE")),
                    "discount_percent": to_decimal(row.get("DISCOUNT_PERCENT")),
                })

        if current:
            groups.append(current)

        if not groups:
            return Response({"error": "No valid purchase rows found in file"}, status=400)

        # ── PASS 2: resolve + validate each group (PER-ENTRY validation) ──
        resolved_groups = []
        for g in groups:
            row_no = g["row_no"]

            party = suppliers.get(g["party_name_raw"].strip().lower())
            if not party:
                errors.append(f"Row {row_no}: Supplier '{g['party_name_raw']}' not found in this branch")

            terms_val = g["terms_raw"].capitalize()
            if terms_val not in TERMS_CHOICES:
                errors.append(f"Row {row_no}: TERMS must be one of Credit/Cash/Bank (got '{g['terms_raw']}')")

            cash_acc = bank_acc = None
            due_date_filled = not is_empty(g["due_date_raw"])
            cash_filled = bool(g["cash_account_raw"])
            bank_filled = bool(g["bank_account_raw"])

            # ── DUE_DATE / CASH_ACCOUNT / BANK_ACCOUNT — is PARTICULAR
            # entry (row) ke TERMS ke hisaab se sirf wohi ek column bhara
            # hona chahiye. Credit ke liye Due Date COMPULSORY hai.
            if terms_val == "Credit":
                if not due_date_filled:
                    errors.append(f"Row {row_no}: DUE_DATE is required when TERMS = Credit")
                if cash_filled or bank_filled:
                    errors.append(
                        f"Row {row_no}: TERMS = Credit — only DUE_DATE should be filled "
                        f"(remove value from CASH_ACCOUNT/BANK_ACCOUNT)"
                    )
            elif terms_val == "Cash":
                if not cash_filled:
                    errors.append(f"Row {row_no}: CASH_ACCOUNT is required when TERMS = Cash")
                else:
                    cash_acc = cash_accounts.get(g["cash_account_raw"].strip().lower())
                    if not cash_acc:
                        errors.append(f"Row {row_no}: Cash account '{g['cash_account_raw']}' not found for this branch")
                if due_date_filled or bank_filled:
                    errors.append(
                        f"Row {row_no}: TERMS = Cash — only CASH_ACCOUNT should be filled "
                        f"(remove value from DUE_DATE/BANK_ACCOUNT)"
                    )
            elif terms_val == "Bank":
                if not bank_filled:
                    errors.append(f"Row {row_no}: BANK_ACCOUNT is required when TERMS = Bank")
                else:
                    bank_acc = bank_accounts.get(g["bank_account_raw"].strip().lower())
                    if not bank_acc:
                        errors.append(f"Row {row_no}: Bank account '{g['bank_account_raw']}' not found for this branch")
                if due_date_filled or cash_filled:
                    errors.append(
                        f"Row {row_no}: TERMS = Bank — only BANK_ACCOUNT should be filled "
                        f"(remove value from DUE_DATE/CASH_ACCOUNT)"
                    )

            if not g["items"]:
                errors.append(f"Row {row_no}: This purchase entry has no items")

            resolved_items = []
            for it in g["items"]:
                variant = variant_map.get(it["item_raw"].strip().lower())
                if not variant:
                    errors.append(f"Row {it['row_no']}: Item/variant '{it['item_raw']}' not found")
                    continue
                if it["qty"] <= 0:
                    errors.append(f"Row {it['row_no']}: QTY must be greater than 0")
                    continue
                price = it["price"] if it["price"] > 0 else to_decimal(variant.purchasePrice)
                if price <= 0:
                    errors.append(f"Row {it['row_no']}: PRICE must be greater than 0")
                    continue
                resolved_items.append({
                    "item_obj": variant.item,
                    "variant_obj": variant,
                    "qty": it["qty"],
                    "price": price,
                    "discount_percent": it["discount_percent"],
                })

            resolved_groups.append({
                "row_no": row_no,
                "party": party,
                "date_val": parse_date(g["date_raw"], default=datetime.now().date()),
                "purchasebill_no": g["purchasebill_no"],
                "terms": terms_val,
                "due_date": parse_date(g["due_date_raw"]) if terms_val == "Credit" else None,
                "cash_account": cash_acc,
                "bank_account": bank_acc,
                "narration": g["narration"],
                "freight": g["freight"],
                "other_expense": g["other_expense"],
                "round_amount": g["round_amount"],
                "items": resolved_items,
            })

        if errors:
            return Response({"success": False, "errors": errors[:200]}, status=400)

        # ── PASS 3: create everything in one transaction ──
        created_bills = []
        with transaction.atomic():
            for g in resolved_groups:
                purchase = PurchaseMaster.objects.create(
                    branch=branch,
                    partyName=g["party"],
                    billNo=generate_purchase_voucher(branch),
                    purchasebill_no=g["purchasebill_no"],
                    date=g["date_val"],
                    dueDate=g["due_date"],
                    terms=g["terms"],
                    narration=g["narration"],
                    frightcharge=g["freight"],
                    otherexpnse=g["other_expense"],
                    roundamount=g["round_amount"],
                    bank_account=g["bank_account"],
                    case_account=g["cash_account"],
                    created_by=request.user,
                )

                for it in g["items"]:
                    pi = PurchaseItem(
                        purchase=purchase,
                        itemName=it["item_obj"],
                        variant=it["variant_obj"],
                        hsnCode=it["item_obj"].hsnCode or "",
                        quantity=it["qty"],
                        altQuantity=Decimal("0.00"),
                        price=it["price"],
                        per=_unit_symbol(it["item_obj"]),
                        discountPercent=it["discount_percent"],
                        netValue=Decimal("0.00"),
                    )
                    pi.save()  # ✅ EXISTING model logic — GST/discount calc yahin hota hai

                agg = purchase.items.aggregate(b=Sum("basicAmount"), t=Sum("taxAmount"), n=Sum("netValue"))
                total_basic = agg["b"] or Decimal("0.00")
                total_tax = agg["t"] or Decimal("0.00")
                total_net = agg["n"] or Decimal("0.00")
                grand_total = total_net + g["freight"] + g["other_expense"] + g["round_amount"]

                purchase.total_basic = total_basic
                purchase.total_tax = total_tax
                purchase.total_net = total_net
                purchase.grand_total = grand_total
                purchase.save(update_fields=[
                    "total_basic", "total_tax", "total_net", "grand_total"
                ])

                terms_lower = g["terms"].lower()
                if terms_lower == "credit":
                    PurchaseMaster.update_balance(purchase.partyName, purchase.grand_total, "Cr")
                elif terms_lower == "cash" and purchase.case_account and purchase.partyName:
                    self._create_cash_payment(purchase, branch, request)
                elif terms_lower == "bank" and purchase.bank_account and purchase.partyName:
                    self._create_bank_payment(purchase, branch, request)

                created_bills.append({
                    "billNo": purchase.billNo,
                    "party": purchase.partyName.account_name,
                    "grand_total": float(purchase.grand_total),
                    "items_count": len(g["items"]),
                })

        return Response({
            "success": True,
            "message": f"{len(created_bills)} purchase entries created successfully",
            "purchases": created_bills,
        }, status=status.HTTP_201_CREATED)