# pos/services/ai_purchase_excel_generator.py
"""
AI se extract hui invoice data ko EXACT USER TEMPLATE format
(PARTY_NAME, DATE, TERMS, ITEM_VARIANT, ...) me likh kar Excel banata hai.

IMPORTANT: Ye file SIRF Excel generate karta hai — koi import/validation
logic yahan nahi hai. Wo sab pos/views/purchase_excel_views.py
(PurchaseExcelImportView) me already hai aur usse bilkul touch nahi kiya gaya.
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from pos.models.items import itemvariants
from pos.models.account import Account

from pos.views.purchase_excel_views import (
    COLUMNS,
    COL_INDEX,
    TERMS_CHOICES,
    dedupe_preserve_order,
    build_variant_lookup,
    _unit_symbol,
)

UNMATCHED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # halka yellow


def _norm(s):
    return (s or "").strip().lower()


def _parse_ai_date(raw):
    """AI kai formats me date de sakta hai — best-effort parse karke date
    object return karta hai (Excel me date ki tarah dikhega), warna raw text."""
    if not raw:
        return ""
    raw = str(raw).strip()
    formats = ["%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y", "%m/%d/%Y"]
    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return raw


def _to_number(raw, default=""):
    if raw in (None, ""):
        return default
    try:
        s = str(raw).replace(",", "").replace("%", "").strip()
        return float(s)
    except (ValueError, TypeError):
        return default


def _match_party(name, suppliers_by_norm):
    if not name:
        return "", None
    key = _norm(name)
    if key in suppliers_by_norm:
        return suppliers_by_norm[key], True
    for k, original in suppliers_by_norm.items():
        if key in k or k in key:
            return original, True
    return name, False


def _match_item_variant(ai_item, variant_lookup_by_norm):
    """
    Match strategy: 'itemname (size)' exact -> 'itemname' exact -> contains match -> no match.
    Returns (label_to_write, matched_variant_obj_or_None, matched_bool)
    """
    item_name = (ai_item.get("item_name") or "").strip()
    size = (ai_item.get("size") or "").strip()

    if not item_name:
        return "", None, None

    candidates = []
    if size:
        candidates.append(f"{item_name} ({size})")
    candidates.append(item_name)

    for candidate in candidates:
        key = _norm(candidate)
        if key in variant_lookup_by_norm:
            variant_obj, real_label = variant_lookup_by_norm[key]
            return real_label, variant_obj, True

    name_key = _norm(item_name)
    for key, (variant_obj, real_label) in variant_lookup_by_norm.items():
        if name_key in key or key.startswith(name_key):
            return real_label, variant_obj, True

    return item_name, None, False


def generate_ai_filled_excel(branch, extracted: dict):
    """
    Returns: (Workbook, warnings: list[str])
    Structure/dropdowns/hidden-lookup bilkul PurchaseExcelTemplateView jaisa
    hi hai — sirf data AI se prefill hota hai.
    """
    warnings = []

    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Data"

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    master_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    item_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(*(Side(style="thin"),) * 4)

    max_row = 1000
    COLUMN_WIDTHS = {"ITEM_VARIANT": 40, "PARTY_NAME": 28, "NARRATION": 30}

    for idx, col_name in enumerate(COLUMNS, 1):
        required = col_name in ("PARTY_NAME", "DATE", "TERMS", "ITEM_VARIANT", "QTY", "PRICE")
        label = col_name.replace("_", " ") + ("*" if required else "")
        cell = ws.cell(row=1, column=idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
        ws.column_dimensions[get_column_letter(idx)].width = COLUMN_WIDTHS.get(col_name, 20)

    MASTER_ONLY = {
        "PARTY_NAME", "DATE", "PURCHASE_BILL_NO", "TERMS", "DUE_DATE",
        "CASH_ACCOUNT", "BANK_ACCOUNT", "NARRATION",
        "FREIGHT_CHARGE", "OTHER_EXPENSE", "ROUND_AMOUNT",
    }
    for col_name, idx in COL_INDEX.items():
        fill = master_fill if col_name in MASTER_ONLY else item_fill
        for r in range(2, max_row + 1):
            cell = ws.cell(row=r, column=idx)
            cell.fill = fill
            cell.border = thin

    # ===== LOOKUP DATA (bilkul template jaisa — branch filtered, deduped) =====
    suppliers_raw = list(
        Account.objects.filter(group="Supplier", branch=branch)
        .order_by("account_name").values_list("account_name", flat=True)
    )
    cash_accounts_raw = list(
        Account.objects.filter(group="Case In Hand", branch=branch)
        .order_by("account_name").values_list("account_name", flat=True)
    )
    bank_accounts_raw = list(
        Account.objects.filter(group="Bank Account", branch=branch)
        .order_by("account_name").values_list("account_name", flat=True)
    )
    suppliers = dedupe_preserve_order(suppliers_raw)
    cash_accounts = dedupe_preserve_order(cash_accounts_raw)
    bank_accounts = dedupe_preserve_order(bank_accounts_raw)
    suppliers_by_norm = {_norm(s): s for s in suppliers}

    variants_qs = (
        itemvariants.objects.filter(item__branch=branch)
        .select_related("item", "item__unit")
        .distinct()
    )
    variant_lookup = build_variant_lookup(variants_qs)

    seen_item_labels = set()
    item_lookup_rows = []
    variant_lookup_by_norm = {}
    for v, unique_label in variant_lookup:
        key = unique_label.strip().lower()
        if key in seen_item_labels:
            continue
        seen_item_labels.add(key)
        variant_lookup_by_norm[key] = (v, unique_label)
        it = v.item
        item_lookup_rows.append((
            unique_label, it.hsnCode or "", _unit_symbol(it),
            (it.taxSlab or "0").replace("%", ""), v.purchasePrice or 0,
        ))
    item_lookup_rows.sort(key=lambda r: r[0].lower())

    hidden_start = len(COLUMNS) + 2
    PARTY_COL = hidden_start
    CASH_COL = hidden_start + 1
    BANK_COL = hidden_start + 2
    ITEM_LOOKUP_START = hidden_start + 3

    for r, name in enumerate(suppliers, start=2):
        ws.cell(row=r, column=PARTY_COL, value=name)
    for r, name in enumerate(cash_accounts, start=2):
        ws.cell(row=r, column=CASH_COL, value=name)
    for r, name in enumerate(bank_accounts, start=2):
        ws.cell(row=r, column=BANK_COL, value=name)
    for r, row_data in enumerate(item_lookup_rows, start=2):
        for offset, val in enumerate(row_data):
            ws.cell(row=r, column=ITEM_LOOKUP_START + offset, value=val)

    for c in (PARTY_COL, CASH_COL, BANK_COL, ITEM_LOOKUP_START, ITEM_LOOKUP_START + 1,
              ITEM_LOOKUP_START + 2, ITEM_LOOKUP_START + 3, ITEM_LOOKUP_START + 4):
        ws.column_dimensions[get_column_letter(c)].hidden = True

    def col_range(col_idx, count):
        letter = get_column_letter(col_idx)
        return f"${letter}$2:${letter}${max(count + 1, 2)}"

    def add_list_validation(target_col_name, formula):
        idx = COL_INDEX[target_col_name]
        letter = get_column_letter(idx)
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{max_row}")

    add_list_validation("PARTY_NAME", col_range(PARTY_COL, len(suppliers)))
    add_list_validation("TERMS", '"Credit,Cash,Bank"')
    add_list_validation("CASH_ACCOUNT", col_range(CASH_COL, len(cash_accounts)))
    add_list_validation("BANK_ACCOUNT", col_range(BANK_COL, len(bank_accounts)))
    add_list_validation("ITEM_VARIANT", col_range(ITEM_LOOKUP_START, len(item_lookup_rows)))

    lookup_range = (
        f"${get_column_letter(ITEM_LOOKUP_START)}$2:"
        f"${get_column_letter(ITEM_LOOKUP_START + 4)}${max(len(item_lookup_rows) + 1, 2)}"
    )
    item_col_letter = get_column_letter(COL_INDEX["ITEM_VARIANT"])

    # ===== AI DATA -> ROWS =====
    supplier = extracted.get("supplier") or {}
    invoice = extracted.get("invoice") or {}
    totals = extracted.get("totals") or {}
    ai_items = extracted.get("items") or []
    terms_raw = (extracted.get("terms") or "").strip().capitalize()

    party_value, party_matched = _match_party(supplier.get("name"), suppliers_by_norm)
    if supplier.get("name") and not party_matched:
        warnings.append(
            f"Supplier '{supplier.get('name')}' bill me mila lekin system ke existing accounts se "
            f"match nahi hua — Excel me likh diya hai (yellow), please PARTY_NAME dropdown se sahi "
            f"account select karo ya naya account banao."
        )
    elif not supplier.get("name"):
        warnings.append("Supplier/Party ka naam bill me clearly nahi mila — PARTY_NAME khud bharna hoga.")

    terms_value = terms_raw if terms_raw in TERMS_CHOICES else ""
    if not terms_value:
        warnings.append("Payment TERMS (Credit/Cash/Bank) bill se pakka nahi ho paya — TERMS column khud select karo.")
    elif terms_value == "Cash":
        warnings.append("TERMS = Cash detect hua — CASH_ACCOUNT column me sahi account dropdown se select karo (ye bill me print nahi hota).")
    elif terms_value == "Bank":
        warnings.append("TERMS = Bank detect hua — BANK_ACCOUNT column me sahi account dropdown se select karo (ye bill me print nahi hota).")

    date_val = _parse_ai_date(invoice.get("date"))
    if not date_val:
        warnings.append("PURCHASE DATE bill me clearly nahi mili — DATE column khud bharo.")
    due_date_val = _parse_ai_date(invoice.get("due_date")) if terms_value == "Credit" else ""
    if terms_value == "Credit" and not due_date_val:
        warnings.append("TERMS = Credit hai lekin DUE_DATE nahi mili — please DUE_DATE khud bharo (zaroori hai).")

    row_num = 2
    for idx, it in enumerate(ai_items):
        is_first = idx == 0
        item_label, matched_variant, item_matched = _match_item_variant(it, variant_lookup_by_norm)

        if it.get("item_name") and not item_matched:
            warnings.append(
                f"Row {row_num}: Item '{it.get('item_name')}' ko existing item/variant se match nahi "
                f"kar paye — AI ne jo bill me padha wahi likh diya hai (yellow highlight), please "
                f"ITEM VARIANT dropdown se sahi item select karo."
            )

        qty_val = _to_number(it.get("quantity"), "")
        if qty_val == "":
            warnings.append(f"Row {row_num}: QUANTITY bill me clearly nahi mili — khud bharo.")

        price_val = _to_number(it.get("purchase_price"), "")
        if price_val == "" and not item_matched:
            warnings.append(f"Row {row_num}: PRICE bill me clearly nahi mila — khud bharo.")

        row_values = {
            "PARTY_NAME": party_value if is_first else "",
            "DATE": date_val if is_first else "",
            "PURCHASE_BILL_NO": invoice.get("bill_no") if is_first else "",
            "TERMS": terms_value if is_first else "",
            "DUE_DATE": due_date_val if is_first else "",
            "CASH_ACCOUNT": "",
            "BANK_ACCOUNT": "",
            "NARRATION": "",
            "FREIGHT_CHARGE": _to_number(totals.get("freight_charge")) if is_first else "",
            "OTHER_EXPENSE": _to_number(totals.get("other_expense")) if is_first else "",
            "ROUND_AMOUNT": _to_number(totals.get("round_off")) if is_first else "",
            "ITEM_VARIANT": item_label,
            "QTY": qty_val,
            "PRICE": price_val,
            "DISCOUNT_PERCENT": _to_number(it.get("discount_percent"), 0),
        }

        for col_name, value in row_values.items():
            idx_col = COL_INDEX[col_name]
            cell = ws.cell(row=row_num, column=idx_col, value=value if value != "" else None)
            if col_name == "ITEM_VARIANT" and not item_matched:
                cell.fill = UNMATCHED_FILL
            if col_name == "PARTY_NAME" and is_first and not party_matched and party_value:
                cell.fill = UNMATCHED_FILL

        if item_matched:
            ws.cell(row=row_num, column=COL_INDEX["HSN_CODE"],
                     value=f'=IFERROR(VLOOKUP(${item_col_letter}{row_num},{lookup_range},2,FALSE),"")')
            ws.cell(row=row_num, column=COL_INDEX["UNIT"],
                     value=f'=IFERROR(VLOOKUP(${item_col_letter}{row_num},{lookup_range},3,FALSE),"")')
            ws.cell(row=row_num, column=COL_INDEX["TAX_PERCENT"],
                     value=f'=IFERROR(VLOOKUP(${item_col_letter}{row_num},{lookup_range},4,FALSE),"")')
            if price_val == "":
                ws.cell(row=row_num, column=COL_INDEX["PRICE"],
                         value=f'=IFERROR(VLOOKUP(${item_col_letter}{row_num},{lookup_range},5,FALSE),"")')
        else:
            ws.cell(row=row_num, column=COL_INDEX["HSN_CODE"], value=it.get("hsn_code") or "")
            ws.cell(row=row_num, column=COL_INDEX["UNIT"], value=it.get("unit") or "")
            ws.cell(row=row_num, column=COL_INDEX["TAX_PERCENT"], value=_to_number(it.get("gst_rate"), ""))

        row_num += 1

    if not ai_items:
        warnings.append("AI ko bill me koi item row detect nahi hui — items manually Excel me bharo.")

    # Baaki khaali rows me bhi formula rakho, taaki user manually aur items add kar sake
    for r in range(row_num, min(row_num + 20, max_row + 1)):
        ws.cell(row=r, column=COL_INDEX["HSN_CODE"],
                 value=f'=IFERROR(VLOOKUP(${item_col_letter}{r},{lookup_range},2,FALSE),"")')
        ws.cell(row=r, column=COL_INDEX["UNIT"],
                 value=f'=IFERROR(VLOOKUP(${item_col_letter}{r},{lookup_range},3,FALSE),"")')
        ws.cell(row=r, column=COL_INDEX["TAX_PERCENT"],
                 value=f'=IFERROR(VLOOKUP(${item_col_letter}{r},{lookup_range},4,FALSE),"")')
        ws.cell(row=r, column=COL_INDEX["PRICE"],
                 value=f'=IFERROR(VLOOKUP(${item_col_letter}{r},{lookup_range},5,FALSE),"")')

    # ===== INSTRUCTIONS SHEET =====
    ws_instr = wb.create_sheet("Instructions")
    ws_instr.column_dimensions["A"].width = 100
    lines = [
        "AI-GENERATED PURCHASE EXCEL — INSTRUCTIONS",
        "",
        "Is Excel ka data AI ne aapke upload kiye bill se padh kar bhara hai.",
        "AI ne SIRF data padha hai — koi Purchase abhi create NAHI hui hai.",
        "",
        "1. Yellow highlighted cells wo hain jinke liye AI ko system me exact match nahi mila —",
        "   dropdown se sahi value select karo, ya value ko manually correct karo.",
        "2. PARTY_NAME/DATE/TERMS/etc sirf PEHLI item row me bharte hain — baaki rows me khaali chhodo.",
        "3. TERMS = Credit/Cash/Bank — jo select karoge uske hisaab se DUE_DATE ya CASH_ACCOUNT ya",
        "   BANK_ACCOUNT bharna hoga (baaki do khaali rakho).",
        "4. HSN/UNIT/TAX% columns matched items ke liye apne aap system se aa jaate hain (formula) —",
        "   unhe edit mat karo. Unmatched (yellow) items me ye AI ki reading hai, sirf reference ke liye.",
        "5. Excel bharne ke baad 'Import Excel' button se upload karo — final validation aur",
        "   Purchase creation wahi existing rules follow karega jo pehle se system me hai.",
        "6. Import se pehle sab AI-detected values ko ek baar bill se match karke verify zaroor karo.",
    ]
    for i, line in enumerate(lines, 1):
        ws_instr.cell(row=i, column=1, value=line)
    ws_instr.sheet_view.showGridLines = False

    ws.freeze_panes = "A2"
    return wb, warnings